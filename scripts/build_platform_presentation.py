from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path("/home/team-02/lms-course3")
SOURCE = ROOT / "docs" / "aurora-lms-platform-presentation.md"
OUTPUT = ROOT / "docs" / "aurora-lms-platform-presentation.pptx"

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

COLOR_NAVY = RGBColor(15, 49, 93)
COLOR_SKY = RGBColor(15, 124, 159)
COLOR_TEXT = RGBColor(32, 52, 81)
COLOR_MUTED = RGBColor(93, 109, 132)
COLOR_BG = RGBColor(234, 242, 251)
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_LINE = RGBColor(209, 222, 239)


def clean_markdown_text(text: str) -> str:
    cleaned = text.strip()
    cleaned = re.sub(r"\*\*(.*?)\*\*", r"\1", cleaned)
    cleaned = re.sub(r"\*(.*?)\*", r"\1", cleaned)
    cleaned = re.sub(r"`(.*?)`", r"\1", cleaned)
    return cleaned


def parse_slides(markdown: str) -> list[dict[str, object]]:
    title_match = re.search(r"^#\s+(.+)$", markdown, re.MULTILINE)
    intro_lines = []
    if title_match:
        after_title = markdown[title_match.end() :].strip().splitlines()
        for line in after_title:
            if line.startswith("## "):
                break
            if line.strip():
                intro_lines.append(line.strip())

    sections = re.split(r"^##\s+", markdown, flags=re.MULTILINE)
    slides: list[dict[str, object]] = []

    if title_match:
        slides.append(
            {
                "title": title_match.group(1).strip(),
                "subtitle": " ".join(intro_lines).strip(),
                "bullets": [],
                "note": "",
                "is_cover": True,
            }
        )

    for raw in sections[1:]:
        lines = [line.rstrip() for line in raw.splitlines()]
        if not lines:
            continue
        title = lines[0].strip()
        bullets: list[str] = []
        note = ""
        for line in lines[1:]:
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("**Заметка докладчика:**"):
                note = clean_markdown_text(stripped.replace("**Заметка докладчика:**", "").strip())
                continue
            if re.match(r"^\d+\.\s+", stripped):
                bullets.append(clean_markdown_text(re.sub(r"^\d+\.\s+", "", stripped)))
                continue
            if stripped.startswith("- "):
                bullets.append(clean_markdown_text(stripped[2:].strip()))
                continue
            if not bullets and not note:
                bullets.append(clean_markdown_text(stripped))
        slides.append(
            {
                "title": clean_markdown_text(title),
                "subtitle": "",
                "bullets": bullets,
                "note": note,
                "is_cover": False,
            }
        )
    return slides


def add_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text,
    font_size=20,
    bold=False,
    color=COLOR_TEXT,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.space_after = Pt(2)
    run = p.runs[0]
    run.font.name = "Arial"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def build_presentation(slides_data: list[dict[str, object]]):
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    blank = prs.slide_layouts[6]

    for index, slide_data in enumerate(slides_data, start=1):
        slide = prs.slides.add_slide(blank)
        is_cover = bool(slide_data["is_cover"])

        add_rect(slide, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT, COLOR_BG)
        add_rect(slide, Inches(0.35), Inches(0.35), Inches(12.63), Inches(6.8), COLOR_WHITE, COLOR_LINE)
        add_rect(slide, Inches(0.35), Inches(0.35), Inches(12.63), Inches(0.24), COLOR_NAVY)

        if is_cover:
            add_rect(slide, Inches(0.75), Inches(1.0), Inches(5.0), Inches(4.7), COLOR_NAVY)
            add_textbox(
                slide,
                Inches(1.1),
                Inches(1.45),
                Inches(4.2),
                Inches(2.2),
                str(slide_data["title"]),
                font_size=26,
                bold=True,
                color=COLOR_WHITE,
            )
            add_textbox(
                slide,
                Inches(1.1),
                Inches(3.9),
                Inches(4.1),
                Inches(1.0),
                str(slide_data["subtitle"]),
                font_size=14,
                color=RGBColor(220, 232, 246),
            )
            add_textbox(
                slide,
                Inches(6.2),
                Inches(1.4),
                Inches(5.0),
                Inches(0.4),
                "AURORA LMS",
                font_size=15,
                bold=True,
                color=COLOR_SKY,
            )
            cover_points = [
                "Корпоративное обучение в одном контуре",
                "Назначения, тесты, отчеты и роли",
                "Подходит для внутренних команд и внешних клиентов",
            ]
            top = 2.1
            for point in cover_points:
                add_rect(slide, Inches(6.15), Inches(top), Inches(0.18), Inches(0.18), COLOR_SKY)
                add_textbox(slide, Inches(6.45), Inches(top - 0.04), Inches(5.2), Inches(0.45), point, font_size=15)
                top += 0.62
        else:
            add_textbox(
                slide,
                Inches(0.9),
                Inches(0.98),
                Inches(8.4),
                Inches(0.7),
                str(slide_data["title"]),
                font_size=21,
                bold=True,
                color=COLOR_TEXT,
            )

            bullets = [str(item) for item in slide_data["bullets"]]
            bullet_box = slide.shapes.add_textbox(Inches(0.95), Inches(1.78), Inches(7.15), Inches(4.7))
            tf = bullet_box.text_frame
            tf.clear()
            tf.word_wrap = True
            tf.margin_left = 0
            tf.margin_right = 0
            tf.margin_top = 0
            tf.margin_bottom = 0
            for i, bullet in enumerate(bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = f"• {bullet}"
                p.level = 0
                p.alignment = PP_ALIGN.LEFT
                p.space_after = Pt(6)
                run = p.runs[0]
                run.font.name = "Arial"
                run.font.size = Pt(16.5)
                run.font.color.rgb = COLOR_TEXT

            add_rect(slide, Inches(8.75), Inches(1.86), Inches(3.45), Inches(3.1), RGBColor(245, 248, 252), COLOR_LINE)
            add_textbox(
                slide,
                Inches(9.02),
                Inches(2.04),
                Inches(2.55),
                Inches(0.4),
                "Главное",
                font_size=12,
                bold=True,
                color=COLOR_SKY,
            )
            add_textbox(
                slide,
                Inches(9.02),
                Inches(2.36),
                Inches(2.78),
                Inches(2.15),
                str(slide_data["note"]),
                font_size=12,
                color=COLOR_MUTED,
            )

        add_textbox(
            slide,
            Inches(11.75),
            Inches(6.82),
            Inches(0.45),
            Inches(0.25),
            str(index),
            font_size=10,
            color=COLOR_MUTED,
            align=PP_ALIGN.RIGHT,
        )

    prs.save(OUTPUT)


def main():
    slides = parse_slides(SOURCE.read_text(encoding="utf-8"))
    build_presentation(slides)


if __name__ == "__main__":
    main()
