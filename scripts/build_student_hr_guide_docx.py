from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT = Path("/home/team-02/lms-course3")
SOURCE = ROOT / "docs" / "aurora-lms-student-hr-guide.md"
OUTPUT = ROOT / "docs" / "aurora-lms-student-hr-guide.docx"

IMAGE_RE = re.compile(r"^!\[(?P<alt>.*?)\]\((?P<path>.+?)\)\s*$")
ORDERED_RE = re.compile(r"^\d+\.\s+(.*)$")
UNORDERED_RE = re.compile(r"^-\s+(.*)$")


def clean_inline_markdown(text: str) -> str:
    text = text.replace("`", "")
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    return text.strip()


def add_picture(document: Document, image_path: Path, alt_text: str) -> None:
    if not image_path.exists():
      document.add_paragraph(f"[Изображение не найдено: {image_path.name}]")
      return

    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = picture_paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(16.5))

    if alt_text:
        caption = document.add_paragraph(clean_inline_markdown(alt_text))
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(15)
    styles["Heading 3"].font.name = "Arial"
    styles["Heading 3"].font.size = Pt(13)


def build_docx() -> None:
    document = Document()
    configure_document(document)

    for raw_line in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped:
            continue

        image_match = IMAGE_RE.match(stripped)
        if image_match:
            image_path = (SOURCE.parent / image_match.group("path")).resolve()
            add_picture(document, image_path, image_match.group("alt"))
            continue

        if stripped.startswith("### "):
            document.add_heading(clean_inline_markdown(stripped[4:]), level=3)
            continue

        if stripped.startswith("## "):
            document.add_heading(clean_inline_markdown(stripped[3:]), level=2)
            continue

        if stripped.startswith("# "):
            document.add_heading(clean_inline_markdown(stripped[2:]), level=1)
            continue

        ordered_match = ORDERED_RE.match(stripped)
        if ordered_match:
            document.add_paragraph(clean_inline_markdown(ordered_match.group(1)), style="List Number")
            continue

        unordered_match = UNORDERED_RE.match(stripped)
        if unordered_match:
            document.add_paragraph(clean_inline_markdown(unordered_match.group(1)), style="List Bullet")
            continue

        document.add_paragraph(clean_inline_markdown(stripped))

    document.save(OUTPUT)


if __name__ == "__main__":
    build_docx()
